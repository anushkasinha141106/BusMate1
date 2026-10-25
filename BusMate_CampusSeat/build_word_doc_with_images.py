import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def build_docx_with_diagram_images():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    PRIMARY_COLOR = RGBColor(7, 89, 133)     # #075985
    SECONDARY_COLOR = RGBColor(15, 118, 110)  # #0F766E
    DARK_TEXT = RGBColor(23, 32, 51)         # #172033

    # Title Banner
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("SmartBus Architecture & Design Diagrams")
    r_t.font.name = "Arial"
    r_t.font.size = Pt(24)
    r_t.font.bold = True
    r_t.font.color.rgb = PRIMARY_COLOR

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s = p_sub.add_run("Intelligent Demand Prediction, Overcrowding Explanation, and Fair Capacity Allocation System\nPublic Inter-City Transit — Pune, India")
    r_s.font.size = Pt(11)
    r_s.font.italic = True
    r_s.font.color.rgb = SECONDARY_COLOR

    doc.add_paragraph()

    # 1. Overview & Research Positioning
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Executive Overview & Research Positioning")
    r_h1.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph(
        "SmartBus connects Machine Learning, Software Engineering, and Human-Computer Interaction (HCI) into a unified, accessible public transit system. The framework dynamically predicts trip overcrowding risk, generates Explainable AI (SHAP) feature attributions, and calibrates a fair capacity-allocation policy under hard bus capacity limits."
    )

    # Callout Box: Research Contribution
    table_callout = doc.add_table(rows=1, cols=1)
    table_callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_callout.cell(0, 0)
    set_cell_background(cell, "F0F9FF")
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.left_indent = Inches(0.1)
    cp.paragraph_format.right_indent = Inches(0.1)
    run_c = cp.add_run("SmartBus Original Contribution:\n")
    run_c.font.bold = True
    run_c.font.color.rgb = PRIMARY_COLOR
    run_c2 = cp.add_run(
        "\"SmartBus integrates overcrowding prediction and explainable ML with counterfactual fairness analysis to dynamically evaluate and calibrate capacity-allocation policies under hard transportation constraints.\""
    )
    run_c2.font.italic = True
    run_c2.font.color.rgb = DARK_TEXT

    doc.add_paragraph()

    # 2. UML Class Diagram Section with PNG Image
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. UML Class Diagram")
    r_h2.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph(
        "The class diagram specifies the entity data models, modular service abstractions, state context, and React UI components forming the SmartBus application."
    )

    if os.path.exists("class_diagram.png"):
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture("class_diagram.png", width=Inches(6.5))

    doc.add_paragraph()

    # 3. Sequence Diagram 1 with PNG Image
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Sequence Diagram 1: Passenger Seat Allocation & Explanation Flow")
    r_h3.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph(
        "This sequence traces a passenger inspecting overcrowding SHAP explanations (C14), submitting a seat reservation request under hard capacity limits (N <= 50), and receiving a transparent outcome with alternative departure recommendations:"
    )

    if os.path.exists("sequence_diagram_passenger_flow.png"):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture("sequence_diagram_passenger_flow.png", width=Inches(6.5))

    doc.add_paragraph()

    # 4. Sequence Diagram 2 with PNG Image
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Sequence Diagram 2: Researcher Counterfactual Sensitivity & Policy Calibration")
    r_h4.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph(
        "This sequence illustrates an Administrator / Researcher conducting single-attribute perturbation experiments in the Counterfactual Lab (C11), running multi-policy trade-off simulations, and selecting a Pareto-calibrated policy vector:"
    )

    if os.path.exists("sequence_diagram_counterfactual_flow.png"):
        p_img3 = doc.add_paragraph()
        p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img3.add_run().add_picture("sequence_diagram_counterfactual_flow.png", width=Inches(6.5))

    output_file = "SmartBus_Architecture_With_Diagram_Images.docx"
    doc.save(output_file)
    print(f"Word document with embedded PNG diagram images created successfully: {output_file}")

if __name__ == "__main__":
    build_docx_with_diagram_images()
