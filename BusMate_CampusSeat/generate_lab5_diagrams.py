import base64
import urllib.request
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def render_mermaid_to_png(mermaid_code, output_png_path):
    encoded_bytes = base64.b64encode(mermaid_code.encode('utf-8'))
    encoded_str = encoded_bytes.decode('utf-8')
    
    url = f"https://mermaid.ink/img/{encoded_str}?bgColor=FFFFFF"
    
    headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
    }
    
    print(f"Fetching diagram image from mermaid.ink for {output_png_path}...")
    req = urllib.request.Request(url, headers=headers)
    
    try:
        with urllib.request.urlopen(req, timeout=30) as response:
            with open(output_png_path, 'wb') as f:
                f.write(response.read())
        print(f"Successfully saved PNG image: {output_png_path}")
        return True
    except Exception as e:
        print(f"Error fetching from mermaid.ink: {e}")
        return False

# ==================================================
# 1. CLASS DIAGRAM MERMAID (Lab 5 Strict Notation)
# ==================================================
CLASS_DIAGRAM_MMD = """classDiagram
    direction LR
    class Bus {
        +id: String
        +routeNumber: String
        +origin: String
        +destination: String
        +departureTime: String
        +capacity: int
        +currentOccupancy: int
        +predictedDemand: int
        +overcrowdingProbability: float
        +riskLevel: String
        +fare: float
    }

    class PassengerRequest {
        +passengerId: String
        +constraintLevel: int
        +scheduleDisruption: int
        +previousDenials: int
        +preferenceMatch: float
        +isAccessibilityRequired: boolean
    }

    class AllocationService {
        +runFairAllocation(requests: PassengerRequest[], capacity: int, weights: float[]): AllocationResult
    }

    class AllocationResult {
        +totalRequests: int
        +capacity: int
        +allocatedCount: int
        +waitlistedCount: int
        +results: Array
    }

    class CounterfactualService {
        +runSingleCounterfactual(passengerId: String, field: String, value: number): CounterfactualResult
        +runBatchCounterfactualEval(requests: PassengerRequest[]): CounterfactualResult[]
    }

    Bus "1" -- "0..*" PassengerRequest : receives requests for >
    PassengerRequest -- AllocationService : processes >
    AllocationService -- AllocationResult : produces >
    CounterfactualService -- AllocationService : evaluates using >
"""

# ==================================================
# 2. SEQUENCE DIAGRAM 1: PASSENGER SEAT ALLOCATION & SHAP EXPLAINABILITY
# ==================================================
SEQ_DIAGRAM_1_MMD = """sequenceDiagram
    autonumber
    actor Passenger
    participant HomePage as HomePage / BusCard
    participant ExplanationModal
    participant SeatRequestModal
    participant XaiService
    participant AllocationService
    participant LocalStorage

    Passenger->>HomePage: Click "Why is this crowded?" on Bus SB-04
    HomePage->>ExplanationModal: open modal(bus = SB-04)
    ExplanationModal->>XaiService: getSHAPExplanations("SB-04")
    XaiService-->>ExplanationModal: SHAP values (+0.31 Friday, +0.18 Event)
    ExplanationModal-->>Passenger: Simple passenger explanation + technical SHAP chart
    Passenger->>ExplanationModal: Close ExplanationModal
    Passenger->>HomePage: Click "Request Seat" on Bus SB-04
    HomePage->>SeatRequestModal: Open SeatRequestModal for SB-04
    Passenger->>SeatRequestModal: Enter (Urgency=High, Previous Denials=1, Accessibility=Yes)
    Passenger->>SeatRequestModal: Click "Submit & Evaluate Allocation"
    SeatRequestModal->>AllocationService: runFairAllocation(requests, capacity=50, weights)
    AllocationService-->>SeatRequestModal: AllocationResult (Status=Allocated, Seat=Seat-04)
    SeatRequestModal->>LocalStorage: setItem("smartbus_my_trips", bookingData)
    SeatRequestModal-->>Passenger: UI displays "Seat Reserved: Seat-04" + transparent reason
"""

# ==================================================
# 3. SEQUENCE DIAGRAM 2: COUNTERFACTUAL ANALYSIS & POLICY CALIBRATION
# ==================================================
SEQ_DIAGRAM_2_MMD = """sequenceDiagram
    autonumber
    actor Admin as Admin / Researcher
    participant CounterfactualLabPage
    participant CounterfactualService
    participant PolicyCalibrationPage
    participant CalibrationService
    participant AllocationService
    participant LocalStorage

    Admin->>CounterfactualLabPage: Select Pass-014 (Baseline = Waitlisted)
    Admin->>CounterfactualLabPage: Perturb Previous Denials: 0 -> 3
    CounterfactualLabPage->>CounterfactualService: runSingleCounterfactual(id="Pass-014", field="previousDenials", val=3)
    CounterfactualService->>AllocationService: Request baseline allocation
    AllocationService-->>CounterfactualService: Baseline Result: Pass-014 = Waitlisted
    CounterfactualService->>AllocationService: Request perturbed allocation
    AllocationService-->>CounterfactualService: Perturbed Result: Pass-014 = Allocated
    CounterfactualService-->>CounterfactualLabPage: OutcomeChanged = True
    CounterfactualLabPage-->>Admin: Display Outcome Flip Badge + Sensitivity Explanation

    Admin->>PolicyCalibrationPage: Navigate to Policy Calibration Page
    Admin->>PolicyCalibrationPage: Adjust weights (w1=0.35, w2=0.25, w3=0.25, w4=0.15)
    PolicyCalibrationPage->>CalibrationService: evaluateCustomPolicyWeights(weights)
    CalibrationService->>AllocationService: runFairAllocation(requests, capacity=50, weights)
    AllocationService-->>CalibrationService: Return metrics (Fairness=91%, Efficiency=89%)
    CalibrationService-->>PolicyCalibrationPage: Return trade-off / scatter data
    Admin->>PolicyCalibrationPage: Click "Select Recommended Calibrated Policy"
    PolicyCalibrationPage->>LocalStorage: setItem("smartbus_custom_weights", recommendedWeights)
    PolicyCalibrationPage-->>Admin: UI displays "Selected Pareto Knee-Point Policy"
"""

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def generate_all():
    img_class = "class_diagram.png"
    img_seq1 = "sequence_diagram_passenger_flow.png"
    img_seq2 = "sequence_diagram_counterfactual_flow.png"

    render_mermaid_to_png(CLASS_DIAGRAM_MMD, img_class)
    render_mermaid_to_png(SEQ_DIAGRAM_1_MMD, img_seq1)
    render_mermaid_to_png(SEQ_DIAGRAM_2_MMD, img_seq2)

    # Build Word Document
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    PRIMARY_COLOR = RGBColor(7, 89, 133)
    SECONDARY_COLOR = RGBColor(15, 118, 110)
    DARK_TEXT = RGBColor(23, 32, 51)

    # Title Banner
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("SmartBus — Software Engineering Lab 5 Submission")
    r_t.font.name = "Arial"
    r_t.font.size = Pt(22)
    r_t.font.bold = True
    r_t.font.color.rgb = PRIMARY_COLOR

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s = p_sub.add_run("UML Class Diagram & UML Sequence Diagrams\nCorrected, Formatted, and Submission-Ready")
    r_s.font.size = Pt(11)
    r_s.font.italic = True
    r_s.font.color.rgb = SECONDARY_COLOR

    doc.add_paragraph()

    # 1. Class Diagram
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. UML Class Diagram — SmartBus Allocation System")
    r_h1.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph("Correct UML 3-compartment notation with typed attributes (+ attributeName: Type) and method signatures. Establishes explicit associations without unneeded inheritance or floating classes:")

    if os.path.exists(img_class):
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture(img_class, width=Inches(6.5))

    doc.add_paragraph()

    # 2. Sequence Diagram 1
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. UML Sequence Diagram — Passenger Seat Allocation & Explainability")
    r_h2.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph("Traces passenger interaction with SHAP explainability (C14) and fair seat allocation under hard capacity bounds (N <= 50):")

    if os.path.exists(img_seq1):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture(img_seq1, width=Inches(6.5))

    doc.add_paragraph()

    # 3. Sequence Diagram 2
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. UML Sequence Diagram — Counterfactual Analysis & Policy Calibration")
    r_h3.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph("Traces single-attribute perturbation testing in CounterfactualLabPage (C11) and multi-policy Pareto trade-off calibration:")

    if os.path.exists(img_seq2):
        p_img3 = doc.add_paragraph()
        p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img3.add_run().add_picture(img_seq2, width=Inches(6.5))

    word_file1 = "SmartBus_Lab5_UML_Submission.docx"
    word_file2 = "SmartBus_Architecture_With_Diagram_Images.docx"
    doc.save(word_file1)
    doc.save(word_file2)
    print(f"Updated Word documents: {word_file1} and {word_file2}")

if __name__ == "__main__":
    generate_all()
