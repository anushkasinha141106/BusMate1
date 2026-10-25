import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def create_smartbus_docx():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles & Colors
    PRIMARY_COLOR = RGBColor(7, 89, 133)     # #075985
    SECONDARY_COLOR = RGBColor(15, 118, 110)  # #0F766E
    DARK_TEXT = RGBColor(23, 32, 51)         # #172033
    GRAY_TEXT = RGBColor(71, 85, 105)        # #475569

    # Title Banner
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("SmartBus Architecture & Design Diagrams")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Intelligent Demand Prediction, Overcrowding Explanation, and Fair Capacity Allocation System\nPublic Inter-City Transit — Pune, India")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = SECONDARY_COLOR

    doc.add_paragraph() # Spacer

    # 1. Project Overview & Research Positioning
    h1 = doc.add_heading(level=1)
    run_h1 = h1.add_run("1. Executive Overview & Research Positioning")
    run_h1.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph(
        "SmartBus connects Machine Learning, Software Engineering, and Human-Computer Interaction (HCI) into a unified, accessible public transit system. The framework dynamically predicts trip overcrowding risk, generates Explainable AI (SHAP) feature attributions, and calibrates a fair capacity-allocation policy under hard bus capacity limits."
    )

    # Callout Box: Research Contribution
    table_callout = doc.add_table(rows=1, cols=1)
    table_callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_callout.cell(0, 0)
    set_cell_background(cell, "F0F9FF") # Light blue fill
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.left_indent = Inches(0.1)
    cp.paragraph_format.right_indent = Inches(0.1)
    run_c = cp.add_run("SmartBus Original Contribution:\n")
    run_c.font.bold = True
    run_c.font.color.rgb = PRIMARY_COLOR
    run_c2 = cp.add_run(
        "\"SmartBus integrates overcrowding prediction and explainable ML with counterfactual fairness analysis to dynamically evaluate and calibrate capacity-allocation policies under hard transportation constraints.\""
    )
    run_c2.font.italic = True
    run_c2.font.color.rgb = DARK_TEXT

    doc.add_paragraph() # Spacer

    doc.add_paragraph("Literature Foundations:")
    p1 = doc.add_paragraph(style='List Bullet')
    r1 = p1.add_run("C14: ")
    r1.bold = True
    p1.add_run("Overcrowding prediction with Explainable AI (SHAP feature attribution).")

    p2 = doc.add_paragraph(style='List Bullet')
    r2 = p2.add_run("C11: ")
    r2.bold = True
    p2.add_run("Counterfactual explainable fairness and analysis of the fairness–utility trade-off.")

    # 2. 9-Stage System Pipeline
    h2 = doc.add_heading(level=1)
    run_h2 = h2.add_run("2. Main System Pipeline Architecture (9 Stages)")
    run_h2.font.color.rgb = PRIMARY_COLOR

    stages = [
        ("Stage 1", "Historical Transport Data", "Ingests passenger boarding logs, route schedules, day vectors, campus events, and weather."),
        ("Stage 2", "Demand Prediction", "Applies regression ML model predicting total passenger demand volume (e.g. 65 passengers)."),
        ("Stage 3", "Overcrowding Risk", "Multi-factor classifier calculates overcrowding probability percentage (e.g. 84%)."),
        ("Stage 4", "XAI / SHAP Explanation", "Computes additive SHAP feature contributions (+0.31 Friday peak, +0.18 event end)."),
        ("Stage 5", "Fair Capacity Allocation", "Linear priority scoring under hard capacity limits (N <= 50) with priority overrides."),
        ("Stage 6", "Counterfactual Experiments", "Single-attribute perturbation tests (x_i -> x_i'; denials 0 -> 3) evaluating fairness sensitivity."),
        ("Stage 7", "Fairness–Efficiency Calibration", "Simulates candidate weight vectors locating Pareto knee-point balancing group fairness & utilization."),
        ("Stage 8", "Final Seat Allocation", "Confirms seat reservations, assigns seat numbers, and routes waitlists to alternative departures."),
        ("Stage 9", "Passenger Explanation & Feedback", "Displays clear decision reasons and collects user ratings for continuous policy tuning.")
    ]

    pipe_table = doc.add_table(rows=1, cols=3)
    pipe_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = pipe_table.rows[0].cells
    hdr_cells[0].text = "Stage #"
    hdr_cells[1].text = "Pipeline Module"
    hdr_cells[2].text = "Technical Functionality"
    for cell in hdr_cells:
        set_cell_background(cell, "075985")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    for stage_num, stage_title, stage_desc in stages:
        row_cells = pipe_table.add_row().cells
        row_cells[0].text = stage_num
        row_cells[1].text = stage_title
        row_cells[2].text = stage_desc
        for i, c in enumerate(row_cells):
            if stage_num.endswith("1") or stage_num.endswith("3") or stage_num.endswith("5") or stage_num.endswith("7") or stage_num.endswith("9"):
                set_cell_background(c, "F8FAFC")

    doc.add_paragraph() # Spacer

    # 3. UML Class Diagram Section
    h3 = doc.add_heading(level=1)
    run_h3 = h3.add_run("3. UML Class Diagram")
    run_h3.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph(
        "The class diagram specifies the entity data models, modular service abstractions, state context, and React UI components forming the SmartBus application."
    )

    # Class Diagram Mermaid Code Box
    class_code = """classDiagram
    direction TB
    class Bus {
        +string id
        +string routeNumber
        +string origin
        +string destination
        +string departureTime
        +number capacity
        +number currentOccupancy
        +number predictedDemand
        +number overcrowdingProbability
        +string riskLevel
        +string fare
    }
    class PassengerRequest {
        +string passengerId
        +number constraintLevel
        +number scheduleDisruption
        +number previousDenials
        +number preferenceMatch
        +boolean isAccessibilityRequired
    }
    class AllocationResult {
        +number totalRequests
        +number capacity
        +number allocatedCount
        +number waitlistedCount
        +Array~PassengerRequest~ results
    }
    class ShapExplanation {
        +string busId
        +number baseValue
        +Array~ShapFeature~ features
    }
    class AllocationService {
        +runFairAllocation(requests, capacity, weights)
    }
    class CounterfactualService {
        +runSingleCounterfactual()
        +runBatchCounterfactualEval()
    }
    class CalibrationService {
        +simulatePolicyTradeoffs()
        +evaluateCustomPolicyWeights()
    }

    AllocationService --> PassengerRequest : processes
    AllocationService --> AllocationResult : generates
    CounterfactualService --> AllocationService : evaluates
    CalibrationService --> AllocationService : tunes weights"""

    p_code1 = doc.add_paragraph()
    r_code1 = p_code1.add_run(class_code)
    r_code1.font.name = "Courier New"
    r_code1.font.size = Pt(9.5)
    r_code1.font.color.rgb = DARK_TEXT

    doc.add_paragraph()

    # 4. Sequence Diagrams Section
    h4 = doc.add_heading(level=1)
    run_h4 = h4.add_run("4. Sequence Diagrams")
    run_h4.font.color.rgb = PRIMARY_COLOR

    # Sequence Diagram A
    doc.add_heading(level=2, text="4.1 Sequence Diagram A: Passenger Seat Allocation & Transparent Explanation")
    doc.add_paragraph(
        "This sequence traces a passenger inspecting overcrowding SHAP explanations (C14), submitting a seat reservation request under hard capacity limits (N <= 50), and receiving a transparent outcome with alternative departure recommendations."
    )

    seq1_text = """1. Passenger -> UI (HomePage/BusCard): Click "Why is this crowded?" on Bus SB-04
2. UI -> ExplanationModal: Open modal(bus = SB-04)
3. ExplanationModal -> XaiService: getShapExplanations("SB-04")
4. XaiService -->> ExplanationModal: Return SHAP values (+0.31 Friday, +0.18 Event)
5. ExplanationModal -->> Passenger: Display Simple Passenger Text & Technical SHAP Bar Chart
6. Passenger -> UI: Click "Request Seat" on Bus SB-04
7. UI -> SeatRequestModal: Open modal(bus = SB-04)
8. Passenger -> SeatRequestModal: Submit Form (Urgency=High, Denials=1, Accessibility=Yes)
9. SeatRequestModal -> AllocationService: runFairAllocation(requests, capacity=50, weights)
   [AllocationService calculates priority scores & applies accessibility override under N <= 50]
10. AllocationService -->> SeatRequestModal: Return AllocationResult (Status: Allocated, Seat: "Seat-04")
11. SeatRequestModal -> LocalStorage: setItem("smartbus_my_trips", bookingData)
12. SeatRequestModal -->> Passenger: Display "Seat Reserved: Seat-04" + Transparent Reason"""

    p_seq1 = doc.add_paragraph()
    r_seq1 = p_seq1.add_run(seq1_text)
    r_seq1.font.name = "Courier New"
    r_seq1.font.size = Pt(9.5)

    doc.add_paragraph()

    # Sequence Diagram B
    doc.add_heading(level=2, text="4.2 Sequence Diagram B: Counterfactual Sensitivity & Policy Calibration Workflow")
    doc.add_paragraph(
        "This sequence illustrates an Administrator / Researcher conducting single-attribute perturbation experiments in the Counterfactual Lab (C11), running multi-policy trade-off simulations, and selecting a Pareto-calibrated policy vector."
    )

    seq2_text = """1. Admin -> CounterfactualLabPage: Select Target Passenger "Pass-014" (Baseline: Waitlisted)
2. Admin -> CounterfactualLabPage: Perturb Attribute: "Previous Denials: 0 -> 3"
3. CounterfactualLabPage -> CounterfactualService: runSingleCounterfactual(id="Pass-014", field="previousDenials", val=3)
4. CounterfactualService -> AllocationService: runFairAllocation(baselineRequests, capacity=50, weights)
5. AllocationService -->> CounterfactualService: Baseline Result (Pass-014 = Waitlisted)
6. CounterfactualService -> AllocationService: runFairAllocation(modifiedRequests, capacity=50, weights)
7. AllocationService -->> CounterfactualService: Perturbed Result (Pass-014 = Allocated)
8. CounterfactualService -->> CounterfactualLabPage: Return Diff (OutcomeChanged = True)
9. CounterfactualLabPage -->> Admin: Display Outcome Flip Badge & Sensitivity Explanation
10. Admin -> PolicyCalibrationPage: Adjust Sliders (w1=0.35, w2=0.25, w3=0.25, w4=0.15)
11. PolicyCalibrationPage -> CalibrationService: evaluateCustomPolicyWeights(weights)
12. CalibrationService -> AllocationService: runFairAllocation(requests, capacity=50, weights)
13. AllocationService -->> CalibrationService: Return Metrics (Fairness: 91%, Efficiency: 89%)
14. CalibrationService -->> PolicyCalibrationPage: Return Trade-off Scatter Data
15. Admin -> PolicyCalibrationPage: Click "Select Recommended Calibrated Policy"
16. PolicyCalibrationPage -> LocalStorage: setItem("smartbus_custom_weights", recommendedWeights)
17. PolicyCalibrationPage -->> Admin: Display "Selected Pareto Knee-Point Policy" Notification"""

    p_seq2 = doc.add_paragraph()
    r_seq2 = p_seq2.add_run(seq2_text)
    r_seq2.font.name = "Courier New"
    r_seq2.font.size = Pt(9.5)

    # Output filename
    output_filename = "SmartBus_Architecture_and_Diagrams.docx"
    doc.save(output_filename)
    print(f"Document successfully created: {output_filename}")

if __name__ == "__main__":
    create_smartbus_docx()
